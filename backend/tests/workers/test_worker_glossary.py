"""
Tests for translate_worker.py Phase 2 extensions:
- Glossary loaded once before batch loop (GLOS-03)
- glossary kwarg passed to translate_batch (not hardcoded None)
- run_post_check called per batch after translation (GLOS-04)
- run_post_check NOT defined in translate_worker.py (import-only from glossary_service)
"""
from __future__ import annotations

import uuid
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
import pytest_asyncio
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.db.models import Base, Glossary, GlossaryTerm, Job, JobStatus, Segment

TEST_DB_URL = "sqlite+aiosqlite:///:memory:"


@pytest_asyncio.fixture(scope="module")
async def engine():
    e = create_async_engine(TEST_DB_URL, echo=False)
    async with e.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    yield e
    await e.dispose()


@pytest_asyncio.fixture
async def session(engine):
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with factory() as sess:
        yield sess
        await sess.rollback()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


async def _make_job_with_glossary(session, glossary_id: str | None = None) -> Job:
    """Create a job with optional glossary_id."""
    job = Job(
        id=str(uuid.uuid4()),
        status=JobStatus.queued,
        source_lang="en",
        target_lang="vi",
        input_format="docx",
        input_path="/tmp/source.docx",
        original_filename="test.docx",
        glossary_id=glossary_id,
    )
    session.add(job)
    await session.commit()
    await session.refresh(job)
    return job


async def _make_glossary_with_terms(session) -> Glossary:
    """Create a glossary with one term."""
    g = Glossary(
        id=str(uuid.uuid4()),
        name="Test Glossary",
        source_lang="en",
        target_lang="vi",
    )
    session.add(g)
    await session.commit()

    term = GlossaryTerm(
        glossary_id=g.id,
        source_term="API",
        target_term="API",
    )
    session.add(term)
    await session.commit()
    return g


# ---------------------------------------------------------------------------
# Structural checks (no runtime needed)
# ---------------------------------------------------------------------------


def test_run_post_check_not_defined_in_worker():
    """run_post_check must NOT be defined in translate_worker.py — import only."""
    import inspect
    import app.workers.translate_worker as worker_module

    # run_post_check should not be defined in the worker module's own source
    worker_source = inspect.getsource(worker_module)
    assert "def run_post_check" not in worker_source, (
        "run_post_check must be imported from glossary_service, not defined in translate_worker"
    )


def test_glossary_none_not_hardcoded_in_worker():
    """The hardcoded glossary=None Phase 1 line should be replaced with glossary=glossary."""
    import inspect
    import app.workers.translate_worker as worker_module

    worker_source = inspect.getsource(worker_module)
    # The Phase 1 hardcoded None should be gone
    assert "glossary=None,  # Phase 1: no glossary" not in worker_source, (
        "Hardcoded glossary=None (Phase 1) must be replaced with glossary=glossary"
    )


def test_worker_imports_glossary_service_functions():
    """translate_worker.py must import load_glossary_terms_for_job and run_post_check."""
    import inspect
    import app.workers.translate_worker as worker_module

    worker_source = inspect.getsource(worker_module)
    assert "load_glossary_terms_for_job" in worker_source, (
        "Worker must import load_glossary_terms_for_job from glossary_service"
    )
    assert "run_post_check" in worker_source, (
        "Worker must import and call run_post_check from glossary_service"
    )


# ---------------------------------------------------------------------------
# Runtime: glossary injection into translate_batch
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_worker_passes_glossary_to_translate_batch(session, tmp_path):
    """Worker passes real glossary dict (not None) when job has glossary_id."""
    from docx import Document as DocxDocument
    doc = DocxDocument()
    doc.add_paragraph("API call")
    input_path = str(tmp_path / "source.docx")
    doc.save(input_path)

    g = await _make_glossary_with_terms(session)
    job = await _make_job_with_glossary(session, glossary_id=g.id)
    job.input_path = input_path
    job.input_format = "docx"
    await session.commit()

    translate_calls: list[dict] = []
    post_check_calls: list[dict] = []

    async def mock_translate_batch(client, segments, source_lang, target_lang, glossary):
        translate_calls.append({"glossary": glossary, "segments": segments})
        return [f"translated: {s}" for s in segments]

    async def mock_post_check(session, batch_segs, translated_map, glossary, source_lang, target_lang, expansion_thresholds, job_id):
        post_check_calls.append({"glossary": glossary})

    mock_settings = MagicMock()
    mock_settings.data_dir = str(tmp_path)
    mock_settings.token_budget = 3000
    mock_settings.worker_concurrency = 1
    mock_settings.expansion_thresholds_dict = {"en->vi": 1.3}

    mock_redis = AsyncMock()
    mock_redis.publish = AsyncMock()

    ctx = {
        "llm_client": AsyncMock(),
        "redis": mock_redis,
        "settings": mock_settings,
    }

    session_factory = MagicMock()
    session_factory.return_value.__aenter__ = AsyncMock(return_value=session)
    session_factory.return_value.__aexit__ = AsyncMock(return_value=False)
    ctx["session_factory"] = session_factory

    with patch("app.workers.translate_worker.translate_batch", mock_translate_batch):
        with patch("app.workers.translate_worker.run_post_check", mock_post_check):
            from app.workers.translate_worker import _run_translation
            await _run_translation(ctx, session, job.id)

    # Glossary should have been passed (not None) to translate_batch
    assert len(translate_calls) > 0
    passed_glossary = translate_calls[0]["glossary"]
    assert passed_glossary is not None
    assert "API" in passed_glossary, f"Expected 'API' in glossary, got: {passed_glossary}"


@pytest.mark.asyncio
async def test_worker_passes_none_glossary_when_no_glossary(session, tmp_path):
    """Worker passes None glossary when job has no glossary_id."""
    from docx import Document as DocxDocument
    doc = DocxDocument()
    doc.add_paragraph("Hello world")
    input_path = str(tmp_path / "source2.docx")
    doc.save(input_path)

    job = await _make_job_with_glossary(session, glossary_id=None)
    job.input_path = input_path
    job.input_format = "docx"
    await session.commit()

    translate_calls: list[dict] = []

    async def mock_translate_batch(client, segments, source_lang, target_lang, glossary):
        translate_calls.append({"glossary": glossary})
        return [f"translated: {s}" for s in segments]

    async def mock_post_check(session, batch_segs, translated_map, glossary, source_lang, target_lang, expansion_thresholds, job_id):
        pass

    mock_settings = MagicMock()
    mock_settings.data_dir = str(tmp_path)
    mock_settings.token_budget = 3000
    mock_settings.worker_concurrency = 1
    mock_settings.expansion_thresholds_dict = {}

    ctx = {
        "llm_client": AsyncMock(),
        "redis": AsyncMock(),
        "settings": mock_settings,
    }
    ctx["redis"].publish = AsyncMock()

    with patch("app.workers.translate_worker.translate_batch", mock_translate_batch):
        with patch("app.workers.translate_worker.run_post_check", mock_post_check):
            from app.workers.translate_worker import _run_translation
            await _run_translation(ctx, session, job.id)

    assert len(translate_calls) > 0
    assert translate_calls[0]["glossary"] is None


@pytest.mark.asyncio
async def test_worker_calls_run_post_check_per_batch(session, tmp_path):
    """Worker calls run_post_check after each batch completes."""
    from docx import Document as DocxDocument
    doc = DocxDocument()
    doc.add_paragraph("Sentence one")
    doc.add_paragraph("Sentence two")
    input_path = str(tmp_path / "source3.docx")
    doc.save(input_path)

    job = await _make_job_with_glossary(session, glossary_id=None)
    job.input_path = input_path
    job.input_format = "docx"
    await session.commit()

    post_check_calls: list[dict] = []

    async def mock_translate_batch(client, segments, source_lang, target_lang, glossary):
        return [f"translated: {s}" for s in segments]

    async def mock_post_check(session, batch_segs, translated_map, glossary, source_lang, target_lang, expansion_thresholds, job_id):
        post_check_calls.append({
            "batch_size": len(batch_segs),
            "translated_map_keys": list(translated_map.keys()),
        })

    mock_settings = MagicMock()
    mock_settings.data_dir = str(tmp_path)
    mock_settings.token_budget = 3000
    mock_settings.worker_concurrency = 1
    mock_settings.expansion_thresholds_dict = {}

    ctx = {
        "llm_client": AsyncMock(),
        "redis": AsyncMock(),
        "settings": mock_settings,
    }
    ctx["redis"].publish = AsyncMock()

    with patch("app.workers.translate_worker.translate_batch", mock_translate_batch):
        with patch("app.workers.translate_worker.run_post_check", mock_post_check):
            from app.workers.translate_worker import _run_translation
            await _run_translation(ctx, session, job.id)

    # run_post_check must have been called at least once
    assert len(post_check_calls) >= 1
    # Each call's translated_map_keys should be non-empty
    for call in post_check_calls:
        assert len(call["translated_map_keys"]) > 0
