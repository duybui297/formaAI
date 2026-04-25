"""
Unit tests for app.pipeline.pptx.smartart — is_smartart() and extract_smartart_text().

Tests cover:
- Primary enum detection path (returns True for IGX_GRAPHIC)
- Fallback XML URI detection path
- Exception handling (both primary and fallback raise)
- extract_smartart_text happy path and exception fallback
"""
from __future__ import annotations

from unittest.mock import MagicMock, patch, PropertyMock

from app.pipeline.pptx.smartart import (
    SMARTART_URI,
    _DML_NS,
    is_smartart,
    extract_smartart_text,
)


# ---------------------------------------------------------------------------
# is_smartart — primary path (enum check)
# ---------------------------------------------------------------------------

def test_is_smartart_returns_true_for_igx_graphic():
    """Primary path: shape_type == MSO_SHAPE_TYPE.IGX_GRAPHIC → True."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    shape = MagicMock()
    shape.shape_type = MSO_SHAPE_TYPE.IGX_GRAPHIC
    assert is_smartart(shape) is True


def test_is_smartart_returns_false_for_text_box():
    """Primary path: non-SmartArt shape type → False (no XML match either)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    shape = MagicMock()
    shape.shape_type = MSO_SHAPE_TYPE.TEXT_BOX
    # element has no matching graphicData
    shape.element.findall.return_value = []
    assert is_smartart(shape) is False


def test_is_smartart_primary_raises_uses_fallback_xml_match():
    """Primary enum check raises → falls back to XML URI detection → True."""
    shape = MagicMock()
    # Make shape_type access raise
    type(shape).shape_type = PropertyMock(side_effect=Exception("enum error"))

    # Set up XML fallback to match SmartArt URI
    gd = MagicMock()
    gd.get.return_value = SMARTART_URI
    shape.element.findall.return_value = [gd]

    assert is_smartart(shape) is True


def test_is_smartart_both_paths_raise_returns_false():
    """Both primary and fallback raise → returns False (defensive)."""
    shape = MagicMock()
    type(shape).shape_type = PropertyMock(side_effect=Exception("enum error"))
    shape.element.findall.side_effect = Exception("xml error")

    assert is_smartart(shape) is False


def test_is_smartart_xml_fallback_no_match_returns_false():
    """Fallback XML path found but URI doesn't match SmartArt URI → False."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    shape = MagicMock()
    shape.shape_type = MSO_SHAPE_TYPE.TABLE  # not IGX_GRAPHIC

    gd = MagicMock()
    gd.get.return_value = "http://some.other/namespace"  # wrong URI
    shape.element.findall.return_value = [gd]

    assert is_smartart(shape) is False


# ---------------------------------------------------------------------------
# extract_smartart_text
# ---------------------------------------------------------------------------

def test_extract_smartart_text_returns_joined_text():
    """extract_smartart_text joins //a:t text nodes with spaces."""
    shape = MagicMock()

    t1 = MagicMock()
    t1.text = "Hello"
    t2 = MagicMock()
    t2.text = "World"
    shape.element.findall.return_value = [t1, t2]

    result = extract_smartart_text(shape)
    assert result == "Hello World"


def test_extract_smartart_text_skips_empty_nodes():
    """extract_smartart_text skips None and whitespace-only text nodes."""
    shape = MagicMock()

    t1 = MagicMock()
    t1.text = "Content"
    t2 = MagicMock()
    t2.text = None
    t3 = MagicMock()
    t3.text = "   "
    shape.element.findall.return_value = [t1, t2, t3]

    result = extract_smartart_text(shape)
    assert result == "Content"


def test_extract_smartart_text_raises_returns_empty_string():
    """extract_smartart_text returns '' when element.findall raises (defensive)."""
    shape = MagicMock()
    shape.element.findall.side_effect = Exception("xml parse error")

    result = extract_smartart_text(shape)
    assert result == ""
