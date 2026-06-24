"""
Excel/XLSX reassembly — write translated text back into a workbook.

Uses the same walk-order as extract_xlsx_segments() so segments align
with the correct cells.
"""
from __future__ import annotations


def reassemble_xlsx(
    file_path: str,
    segments: list,
    translated_texts: dict[str, str],
) -> str:
    """
    Write translated text back into XLSX cells.

    translated_texts is keyed by segment.id (D-06 sha256 ID).

    Returns the path to the saved output file.
    The traversal order must exactly match extract_xlsx_segments().

    Strategy:
    - Re-open the workbook in write mode.
    - Walk all sheets, rows, and cells in the same order as extraction.
    - For each cell, check if a matching segment ID exists in translated_texts.
    - If found, overwrite the cell value with the translated text.
    - Preserve all other cell formatting (number format, font, fill, alignment).
    - Save to output_path = dir(file_path) / output.xlsx
    """
    import os
    import unicodedata

    from openpyxl import load_workbook

    def _nfc(s: str) -> str:
        return unicodedata.normalize("NFC", s)

    wb = load_workbook(file_path)
    seq = 0
    translated_count = 0

    for sheet in wb.worksheets:
        sheet_name = sheet.title or ""
        for row_idx, row in enumerate(sheet.iter_rows()):
            for col_idx, cell in enumerate(row):
                raw_value = cell.value
                if raw_value is None:
                    continue
                text = _nfc(str(raw_value).strip())
                if not text:
                    continue

                structural_position = (
                    f"sheet.{sheet_name}.row.{row_idx}.col.{col_idx}"
                )

                # Reconstruct segment ID the same way as extraction
                from app.pipeline.segment import make_segment_id

                seg_id = make_segment_id(text, structural_position)
                translated = translated_texts.get(seg_id)
                if translated is not None:
                    cell.value = _nfc(translated)
                    translated_count += 1
                seq += 1

    out_dir = os.path.dirname(file_path)
    out_path = os.path.join(out_dir, "output.xlsx")
    wb.save(out_path)
    wb.close()
    return out_path
