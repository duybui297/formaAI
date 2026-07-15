"""
US-3.6: Lightweight word-count estimation for uploaded documents.

Provides a fast, no-side-effect word-count that doesn't require:
- Creating a job row
- Persisting the file
- Enqueuing any worker task

Used by POST /translations/estimate to show users a credit cost preview
before they commit to starting a translation.
"""
from __future__ import annotations

import re
import tempfile
from pathlib import Path

from app.licensing.plans import calculate_credit_cost


# ---------------------------------------------------------------------------
# Word-count extractors per format
# ---------------------------------------------------------------------------


def _strip_tags(html: str) -> str:
    """Remove XML/HTML tags from text."""
    return re.sub(r"<[^>]+>", "", html)


def _count_words(text: str) -> int:
    """Count whitespace-delimited tokens in text."""
    return len(text.split())


def count_words_docx(content: bytes) -> int:
    """Count words in a DOCX file from raw bytes."""
    try:
        from docx import Document  # noqa: PLC0415
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        try:
            doc = Document(tmp_path)
            # Collect all paragraph text (walk all content like extract_run_segments)
            texts: list[str] = []
            for block in doc.iter_inner_content():
                if hasattr(block, "text"):
                    texts.append(block.text)
            return sum(_count_words(t) for t in texts)
        finally:
            Path(tmp_path).unlink(missing_ok=True)
    except Exception:
        return 0


def count_words_pptx(content: bytes) -> int:
    """Count words in a PPTX file from raw bytes."""
    try:
        from pptx import Presentation as PPTXPresentation  # noqa: PLC0415
        from pptx.util import Inches as PPTXInches  # noqa: PLC0415
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        try:
            prs = PPTXPresentation(tmp_path)
            total = 0
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            total += _count_words(para.text)
            return total
        finally:
            Path(tmp_path).unlink(missing_ok=True)
    except Exception:
        return 0


def count_words_pdf(content: bytes) -> int:
    """Count words in a native PDF from raw bytes."""
    try:
        import pymupdf  # noqa: PLC0415
        doc = pymupdf.open(stream=content, filetype="pdf")
        try:
            total = 0
            for page in doc:
                text = page.get_text("text")
                total += _count_words(text)
            return total
        finally:
            doc.close()
    except Exception:
        return 0


def count_words_xlsx(content: bytes) -> int:
    """Count words in an XLSX file from raw bytes."""
    try:
        from openpyxl import load_workbook  # noqa: PLC0415
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        try:
            wb = load_workbook(tmp_path, data_only=True)
            total = 0
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    for cell in row:
                        if cell is not None:
                            total += _count_words(str(cell))
            return total
        finally:
            Path(tmp_path).unlink(missing_ok=True)
    except Exception:
        return 0


def count_words_in_file(content: bytes, ext: str) -> int:
    """
    Dispatch word-counting to the right handler by file extension.

    Returns 0 on any parse error (estimation fails gracefully).
    """
    match ext.lower():
        case ".docx":
            return count_words_docx(content)
        case ".pptx":
            return count_words_pptx(content)
        case ".pdf":
            return count_words_pdf(content)
        case ".xlsx":
            return count_words_xlsx(content)
        case _:
            return 0


def estimate_credit_cost(
    word_count: int,
    tier_name: str,
    is_scanned: bool,
) -> int:
    """
    Calculate credit cost from word count + tier string.

    tier_name: one of "TRIAL", "PRO", "ENTERPRISE"
    Returns 0 if tier is unrecognized.
    """
    from app.db.models import LicenseTier
    try:
        tier = LicenseTier[tier_name.upper()]
    except KeyError:
        return 0
    return calculate_credit_cost(word_count, tier, is_scanned)
